# d:\programas\DDS\dds-webtools\bdo\services\exportador_service.py
from __future__ import annotations
import io
import re
import math
import zipfile
import datetime as dt
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Set, Tuple

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter as _gcl

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm, inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from bdo.services.constantes import HEADERS_VIZ, MAP_BOL, MAP_PTO

# Paleta de cores profissional
CORES = {
    'primaria': '1F4788',      # Azul escuro profissional
    'secundaria': '4472C4',    # Azul médio
    'destaque': '2E75B6',      # Azul destaque
    'sucesso': '70AD47',       # Verde sucesso
    'alerta': 'FFC000',        # Amarelo/Laranja
    'erro': 'C00000',          # Vermelho
    'neutro_claro': 'F2F2F2',  # Cinza muito claro
    'neutro_medio': 'D9E1F2',  # Azul acinzentado claro
    'branco': 'FFFFFF',        # Branco
    'boletim': 'E7F0FF',       # Azul muito claro
    'ponto': 'E7FBEA',         # Verde muito claro
    'diferenca': 'FFF4E7',     # Laranja muito claro
}

def criar_borda(estilo='thin', cor='CCCCCC'):
    side = Side(style=estilo, color=cor)
    return Border(left=side, right=side, top=side, bottom=side)

def _eh_total(valor) -> bool:
    s = str(valor or "").strip().upper()
    return s.startswith("TOTAIS") or s.startswith("TOTAL")

def aplicar_cabecalho_profissional(ws, titulo, contrato, dt_ini, dt_fim, boletins, df_cols):
    max_col = max(1, len(df_cols))

    # === Linha 1: TÍTULO ===
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    c = ws.cell(row=1, column=1, value=str(titulo) if titulo is not None else "")
    c.font = Font(name='Calibri', size=16, bold=True, color=CORES['branco'])
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.fill = PatternFill(start_color=CORES['primaria'], end_color=CORES['primaria'], fill_type="solid")

    def _rotulo(r, texto):
        ws.row_dimensions[r].height = 22
        r_cell = ws.cell(row=r, column=1, value=texto)
        r_cell.font = Font(name='Calibri', size=11, bold=True, color=CORES['primaria'])
        r_cell.alignment = Alignment(horizontal="left", vertical="center")
        r_cell.fill = PatternFill(start_color=CORES['neutro_claro'], end_color=CORES['neutro_claro'], fill_type="solid")

    def _valor(r, texto):
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=max_col)
        v_cell = ws.cell(row=r, column=2, value=texto)
        v_cell.font = Font(name='Calibri', size=11)
        v_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # === Linha 2: CONTRATO ===
    _rotulo(2, "Contrato:")
    _valor(2, str(contrato) if contrato is not None else "")

    # === Linha 3: PERÍODO ===
    try:
        periodo_txt = f"{dt_ini:%d/%m/%Y} a {dt_fim:%d/%m/%Y}"
    except Exception:
        periodo_txt = ""
    _rotulo(3, "Período de Apuração:")
    _valor(3, periodo_txt)

    # === Linha 4: BOLETINS ===
    boletins_txt = "-" if (boletins is None or str(boletins).strip() == "") else str(boletins)
    _rotulo(4, "Boletins Incluídos:")
    _valor(4, boletins_txt)

    # Linhas 5-6 de espaçamento
    ws.row_dimensions[5].height = 6
    ws.row_dimensions[6].height = 6

def formatar_planilha_profissional(ws, startrow=6):
    max_row = ws.max_row
    max_col = ws.max_column
    
    # Ajusta largura das colunas
    for c in range(1, max_col + 1):
        col_letter = _gcl(c)
        if c == 1:
            ws.column_dimensions[col_letter].width = 22
        else:
            ws.column_dimensions[col_letter].width = 15
    
    # Formata linha de cabeçalho da tabela
    ws.row_dimensions[startrow].height = 30
    for c in range(1, max_col + 1):
        cell = ws.cell(row=startrow, column=c)
        cell.font = Font(name='Calibri', size=10, bold=True, color=CORES['branco'])
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.fill = PatternFill(start_color=CORES['secundaria'], 
                                end_color=CORES['secundaria'], 
                                fill_type="solid")
        cell.border = criar_borda('medium', CORES['primaria'])
    
    # Formata linhas de dados
    for r in range(startrow + 1, max_row + 1):
        ws.row_dimensions[r].height = 18
        primeira_celula = str(ws.cell(row=r, column=1).value or "")
        is_total = _eh_total(primeira_celula)
        
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            
            if is_total:
                cell.font = Font(name='Calibri', size=10, bold=True, color=CORES['primaria'])
                cell.fill = PatternFill(start_color=CORES['neutro_medio'], 
                                       end_color=CORES['neutro_medio'], 
                                       fill_type="solid")
                cell.border = criar_borda('medium', CORES['primaria'])
                if c == 1:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                if (r - startrow) % 2 == 0:
                    cor_fundo = CORES['branco']
                else:
                    cor_fundo = CORES['neutro_claro']
                
                cell.fill = PatternFill(start_color=cor_fundo, 
                                       end_color=cor_fundo, 
                                       fill_type="solid")
                cell.border = criar_borda('thin', 'CCCCCC')
                
                if c == 1:
                    cell.font = Font(name='Calibri', size=9)
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                else:
                    cell.font = Font(name='Calibri', size=9)
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = '#,##0.00'
                        if cell.value < 0:
                            cell.font = Font(name='Calibri', size=9, bold=True, color=CORES['erro'])
                    elif isinstance(cell.value, str) and cell.value.startswith('-'):
                        cell.font = Font(name='Calibri', size=9, bold=True, color=CORES['erro'])
    
    ws.freeze_panes = f"A{startrow + 1}"
    if max_row > startrow:
        ws.auto_filter.ref = f"A{startrow}:{_gcl(max_col)}{max_row}"

def _decimal_to_hhmm(v: float | int | None) -> str:
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return ""
    try:
        neg = v < 0
        v = abs(float(v))
        h = int(v)
        m = int(round((v - h) * 60))
        if m >= 60:
            h += 1
            m -= 60
        s = f"{h:d}:{m:02d}"
        return f"-{s}" if neg else s
    except Exception:
        return str(v)

def _is_negative_float(s: str) -> bool:
    try:
        val = float(str(s).replace(",", "."))
        return val < 0
    except ValueError:
        return False

def _sheet_name_unique(base, used: set) -> str:
    name = re.sub(r"[][:\\/?*]+", "", str(base)).strip() or "Aba"
    name = name[:31]
    orig = name
    i = 2
    while name in used:
        suf = f"_{i}"
        name = orig[: 31 - len(suf)] + suf
        i += 1
    used.add(name)
    return name

def _ult5_contratos(s):
    if s is None:
        return ""
    itens = [p.strip() for p in str(s).split(",") if p.strip()]
    return ", ".join([(i[-5:] if len(i) >= 5 else i) for i in itens])

# ============================================================
# EXPORTADORES INDIVIDUAIS
# ============================================================

def _align_comparison_frames(
    df_b: pd.DataFrame,
    df_p: pd.DataFrame,
    df_d: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    frames = [frame.copy() for frame in (df_b, df_p, df_d)]
    dates = pd.DatetimeIndex([])
    for frame in frames:
        if not frame.empty and "Data" in frame.columns:
            frame["Data"] = pd.to_datetime(
                frame["Data"], errors="coerce"
            ).dt.floor("D")
            dates = dates.union(pd.DatetimeIndex(frame["Data"].dropna()))
    base = pd.DataFrame({"Data": dates.sort_values()})
    aligned = []
    for frame in frames:
        if frame.empty or "Data" not in frame.columns:
            aligned.append(base.copy())
            continue
        frame = frame.drop_duplicates("Data", keep="last")
        aligned.append(base.merge(frame, on="Data", how="left"))
    return tuple(aligned)


def exportar_comparacao_individual_excel(
    df_b: pd.DataFrame,
    df_p: pd.DataFrame,
    df_d: pd.DataFrame,
    employee: str,
    dt_ini: datetime,
    dt_fim: datetime,
    format_hhmm: bool = False
) -> bytes:
    # 1. Alinha dataframes
    df_b, df_p, df_d = _align_comparison_frames(df_b, df_p, df_d)
    headers_horas = [h for h in HEADERS_VIZ if any(h in frame.columns for frame in (df_b, df_p, df_d))]
    max_col = 1 + len(headers_horas) * 3

    # Monta df final em formato numérico para somatório preciso
    df_final = pd.DataFrame()
    df_final["Data"] = df_b["Data"].dt.strftime("%d/%m/%Y")
    
    for h in headers_horas:
        df_final[f"Boletim {h}"] = pd.to_numeric(df_b[h], errors="coerce").fillna(0.0) if h in df_b.columns else 0.0
    for h in headers_horas:
        df_final[f"Ponto {h}"] = pd.to_numeric(df_p[h], errors="coerce").fillna(0.0) if h in df_p.columns else 0.0
    for h in headers_horas:
        df_final[f"Diferença {h}"] = pd.to_numeric(df_d[h], errors="coerce").fillna(0.0) if h in df_d.columns else 0.0

    # Calcula totais
    totais = ["TOTAIS:"]
    for col in df_final.columns[1:]:
        totais.append(df_final[col].sum())

    # Formata para HH:MM ou Decimal
    if format_hhmm:
        for col in df_final.columns[1:]:
            df_final[col] = df_final[col].apply(_decimal_to_hhmm)
        totais = [totais[0]] + [_decimal_to_hhmm(x) for x in totais[1:]]

    wb = Workbook()
    ws = wb.active
    ws.title = "Comparação"

    # Aplica cabeçalho no topo
    aplicar_cabecalho_profissional(
        ws,
        titulo=f"Relatório Comparativo – Funcionário: {employee}",
        contrato="-",
        dt_ini=dt_ini,
        dt_fim=dt_fim,
        boletins="-",
        df_cols=list(df_final.columns)
    )

    # Escreve grupos de colunas (linha 7)
    ws.row_dimensions[7].height = 25
    ws.cell(row=7, column=1, value="")
    
    b_ini, b_fim = 2, 1 + len(headers_horas)
    p_ini, p_fim = b_fim + 1, b_fim + len(headers_horas)
    d_ini, d_fim = p_fim + 1, p_fim + len(headers_horas)
    
    ws.cell(row=7, column=b_ini, value="Boletim")
    ws.merge_cells(start_row=7, start_column=b_ini, end_row=7, end_column=b_fim)
    
    ws.cell(row=7, column=p_ini, value="Ponto")
    ws.merge_cells(start_row=7, start_column=p_ini, end_row=7, end_column=p_fim)
    
    ws.cell(row=7, column=d_ini, value="Diferença")
    ws.merge_cells(start_row=7, start_column=d_ini, end_row=7, end_column=d_fim)

    for c in range(1, max_col + 1):
        cell = ws.cell(row=7, column=c)
        cell.font = Font(name='Calibri', size=11, bold=True, color=CORES['branco'])
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.fill = PatternFill(start_color=CORES['primaria'], end_color=CORES['primaria'], fill_type="solid")
        cell.border = criar_borda('medium', CORES['primaria'])

    # Escreve cabeçalhos das colunas (linha 8)
    ws.row_dimensions[8].height = 35
    ws.cell(row=8, column=1, value="Data")
    col_idx = 2
    for _ in range(3):
        for h in headers_horas:
            ws.cell(row=8, column=col_idx, value=h)
            col_idx += 1

    for c in range(1, max_col + 1):
        cell = ws.cell(row=8, column=c)
        cell.font = Font(name='Calibri', size=10, bold=True, color=CORES['branco'])
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.fill = PatternFill(start_color=CORES['secundaria'], end_color=CORES['secundaria'], fill_type="solid")
        cell.border = criar_borda('medium', CORES['primaria'])

    # Escreve dados (linha 9 em diante)
    current_row = 9
    for _, row in df_final.iterrows():
        ws.row_dimensions[current_row].height = 20
        for c in range(1, max_col + 1):
            val = row.iloc[c - 1]
            cell = ws.cell(row=current_row, column=c, value=val)
            
            # Cores das colunas
            if b_ini <= c <= b_fim:
                cor_fundo = CORES['boletim']
            elif p_ini <= c <= p_fim:
                cor_fundo = CORES['ponto']
            elif d_ini <= c <= d_fim:
                cor_fundo = CORES['diferenca']
            else:
                cor_fundo = CORES['branco']
                
            cell.fill = PatternFill(start_color=cor_fundo, end_color=cor_fundo, fill_type="solid")
            cell.border = criar_borda('thin', 'CCCCCC')
            cell.font = Font(name='Calibri', size=10)
            
            if c == 1:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if not format_hhmm and isinstance(val, (int, float)):
                    cell.number_format = '#,##0.00'
                    if val < 0:
                        cell.font = Font(name='Calibri', size=10, bold=True, color=CORES['erro'])
                elif format_hhmm and str(val).startswith('-'):
                    cell.font = Font(name='Calibri', size=10, bold=True, color=CORES['erro'])
        current_row += 1

    # Escreve Totais
    ws.row_dimensions[current_row].height = 20
    for c in range(1, max_col + 1):
        val = totais[c - 1]
        cell = ws.cell(row=current_row, column=c, value=val)
        
        cell.font = Font(name='Calibri', size=11, bold=True, color=CORES['primaria'])
        cell.fill = PatternFill(start_color=CORES['neutro_medio'], end_color=CORES['neutro_medio'], fill_type="solid")
        cell.border = criar_borda('medium', CORES['primaria'])
        
        if c == 1:
            cell.alignment = Alignment(horizontal="left", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="right", vertical="center")
            if not format_hhmm and isinstance(val, (int, float)):
                cell.number_format = '#,##0.00'
                if val < 0:
                    cell.font = Font(name='Calibri', size=11, bold=True, color=CORES['erro'])
            elif format_hhmm and str(val).startswith('-'):
                cell.font = Font(name='Calibri', size=11, bold=True, color=CORES['erro'])

    ws.column_dimensions["A"].width = 14
    for c in range(2, max_col + 1):
        ws.column_dimensions[_gcl(c)].width = 12
    ws.freeze_panes = "A9"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def exportar_comparacao_individual_pdf(
    df_b: pd.DataFrame,
    df_p: pd.DataFrame,
    df_d: pd.DataFrame,
    employee: str,
    dt_ini: datetime,
    dt_fim: datetime,
    format_hhmm: bool = False
) -> bytes:
    df_b, df_p, df_d = _align_comparison_frames(df_b, df_p, df_d)
    headers_horas = [h for h in HEADERS_VIZ if any(h in frame.columns for frame in (df_b, df_p, df_d))]

    df_final = pd.DataFrame()
    df_final["Data"] = df_b["Data"].dt.strftime("%d/%m/%Y")
    
    for h in headers_horas:
        df_final[f"Boletim {h}"] = pd.to_numeric(df_b[h], errors="coerce").fillna(0.0) if h in df_b.columns else 0.0
    for h in headers_horas:
        df_final[f"Ponto {h}"] = pd.to_numeric(df_p[h], errors="coerce").fillna(0.0) if h in df_p.columns else 0.0
    for h in headers_horas:
        df_final[f"Diferença {h}"] = pd.to_numeric(df_d[h], errors="coerce").fillna(0.0) if h in df_d.columns else 0.0

    totais = ["TOTAIS:"]
    for col in df_final.columns[1:]:
        totais.append(df_final[col].sum())

    if format_hhmm:
        for col in df_final.columns[1:]:
            df_final[col] = df_final[col].apply(_decimal_to_hhmm)
        totais = [totais[0]] + [_decimal_to_hhmm(x) for x in totais[1:]]

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), rightMargin=20, leftMargin=20, topMargin=20, bottomMargin=12)
    styles = getSampleStyleSheet()

    header_cells = ["Data"] + headers_horas * 3
    
    pdf_rows = [header_cells]
    for _, row in df_final.iterrows():
        pdf_row = [row.iloc[0]]
        for val in row.iloc[1:]:
            if isinstance(val, (int, float)):
                pdf_row.append(f"{val:.2f}".replace(".", ","))
            else:
                pdf_row.append(str(val))
        pdf_rows.append(pdf_row)

    pdf_tot = [totais[0]]
    for val in totais[1:]:
        if isinstance(val, (int, float)):
            pdf_tot.append(f"{val:.2f}".replace(".", ","))
        else:
            pdf_tot.append(str(val))
    pdf_rows.append(pdf_tot)

    col_widths = [1.8 * cm] + [1.36 * cm] * (len(header_cells) - 1)
    tabela = Table(pdf_rows, colWidths=col_widths, repeatRows=1)

    estilo = TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#183C5F")),
        ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
    ])

    b_ini, b_fim = 1, len(headers_horas)
    p_ini, p_fim = b_fim + 1, b_fim + len(headers_horas)
    d_ini, d_fim = p_fim + 1, p_fim + len(headers_horas)

    estilo.add("BACKGROUND", (b_ini, 1), (b_fim, -1), colors.HexColor("#B7DFFB"))
    estilo.add("BACKGROUND", (p_ini, 1), (p_fim, -1), colors.HexColor("#C5EDC1"))
    estilo.add("BACKGROUND", (d_ini, 1), (d_fim, -1), colors.HexColor("#FFDAB9"))

    # Destaca negativos
    for r in range(1, len(pdf_rows)):
        for c in range(d_ini + 1, d_fim + 2):
            val_str = pdf_rows[r][c - 1]
            if val_str.startswith("-") or (not format_hhmm and _is_negative_float(val_str)):
                estilo.add("TEXTCOLOR", (c - 1, r), (c - 1, r), colors.red)
                estilo.add("FONTNAME", (c - 1, r), (c - 1, r), "Helvetica-Bold")

    tabela.setStyle(estilo)
    story = [
        Paragraph("<b>Relatório Comparativo – Funcionário</b>", styles["Title"]),
        Paragraph(f"<b>Funcionário:</b> {employee}", styles["Normal"]),
        Paragraph(f"<b>Período:</b> {dt_ini:%d/%m/%Y} a {dt_fim:%d/%m/%Y}", styles["Normal"]),
        Spacer(1, 8),
        tabela,
    ]
    doc.build(story)
    return buf.getvalue()

def exportar_diferencas_individual_excel(
    df_d: pd.DataFrame,
    employee: str,
    dt_ini: datetime,
    dt_fim: datetime,
    format_hhmm: bool = False
) -> bytes:
    df_d = df_d.sort_values("Data").reset_index(drop=True)
    headers_horas = [h for h in HEADERS_VIZ if h in df_d.columns]
    max_col = 1 + len(headers_horas)

    df_final = pd.DataFrame()
    df_final["Data"] = df_d["Data"].dt.strftime("%d/%m/%Y")
    for h in headers_horas:
        df_final[h] = pd.to_numeric(df_d[h], errors="coerce").fillna(0.0)

    totais = ["TOTAIS:"]
    for col in df_final.columns[1:]:
        totais.append(df_final[col].sum())

    if format_hhmm:
        for col in df_final.columns[1:]:
            df_final[col] = df_final[col].apply(_decimal_to_hhmm)
        totais = [totais[0]] + [_decimal_to_hhmm(x) for x in totais[1:]]

    wb = Workbook()
    ws = wb.active
    ws.title = "Diferenças"

    aplicar_cabecalho_profissional(
        ws,
        titulo=f"Relatório – Diferenças (Boletim − Ponto) – Funcionário: {employee}",
        contrato="-",
        dt_ini=dt_ini,
        dt_fim=dt_fim,
        boletins="-",
        df_cols=list(df_final.columns)
    )

    # Linha de grupo (linha 7)
    ws.row_dimensions[7].height = 25
    ws.cell(row=7, column=1, value="")
    ws.cell(row=7, column=2, value="Diferença")
    ws.merge_cells(start_row=7, start_column=2, end_row=7, end_column=max_col)
    
    for c in range(1, max_col + 1):
        cell = ws.cell(row=7, column=c)
        cell.font = Font(name='Calibri', size=11, bold=True, color=CORES['branco'])
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.fill = PatternFill(start_color=CORES['primaria'], end_color=CORES['primaria'], fill_type="solid")
        cell.border = criar_borda('medium', CORES['primaria'])

    # Linha de cabeçalhos (linha 8)
    ws.row_dimensions[8].height = 35
    ws.cell(row=8, column=1, value="Data")
    for j, h in enumerate(headers_horas, start=2):
        ws.cell(row=8, column=j, value=h)
        
    for c in range(1, max_col + 1):
        cell = ws.cell(row=8, column=c)
        cell.font = Font(name='Calibri', size=10, bold=True, color=CORES['branco'])
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.fill = PatternFill(start_color=CORES['secundaria'], end_color=CORES['secundaria'], fill_type="solid")
        cell.border = criar_borda('medium', CORES['primaria'])

    # Dados
    current_row = 9
    for _, row in df_final.iterrows():
        ws.row_dimensions[current_row].height = 20
        for c in range(1, max_col + 1):
            val = row.iloc[c - 1]
            cell = ws.cell(row=current_row, column=c, value=val)
            
            cell.fill = PatternFill(start_color=CORES['diferenca'], end_color=CORES['diferenca'], fill_type="solid")
            cell.border = criar_borda('thin', 'CCCCCC')
            cell.font = Font(name='Calibri', size=10)
            
            if c == 1:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if not format_hhmm and isinstance(val, (int, float)):
                    cell.number_format = '#,##0.00'
                    if val < 0:
                        cell.font = Font(name='Calibri', size=10, bold=True, color=CORES['erro'])
                elif format_hhmm and str(val).startswith('-'):
                    cell.font = Font(name='Calibri', size=10, bold=True, color=CORES['erro'])
        current_row += 1

    # Totais
    ws.row_dimensions[current_row].height = 20
    for c in range(1, max_col + 1):
        val = totais[c - 1]
        cell = ws.cell(row=current_row, column=c, value=val)
        
        cell.font = Font(name='Calibri', size=11, bold=True, color=CORES['primaria'])
        cell.fill = PatternFill(start_color=CORES['neutro_medio'], end_color=CORES['neutro_medio'], fill_type="solid")
        cell.border = criar_borda('medium', CORES['primaria'])
        
        if c == 1:
            cell.alignment = Alignment(horizontal="left", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="right", vertical="center")
            if not format_hhmm and isinstance(val, (int, float)):
                cell.number_format = '#,##0.00'
                if val < 0:
                    cell.font = Font(name='Calibri', size=11, bold=True, color=CORES['erro'])
            elif format_hhmm and str(val).startswith('-'):
                cell.font = Font(name='Calibri', size=11, bold=True, color=CORES['erro'])

    ws.column_dimensions["A"].width = 14
    for c in range(2, max_col + 1):
        ws.column_dimensions[_gcl(c)].width = 12
    ws.freeze_panes = "A9"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def exportar_diferencas_individual_pdf(
    df_d: pd.DataFrame,
    employee: str,
    dt_ini: datetime,
    dt_fim: datetime,
    format_hhmm: bool = False
) -> bytes:
    df_d = df_d.sort_values("Data").reset_index(drop=True)
    headers_horas = [h for h in HEADERS_VIZ if h in df_d.columns]

    df_final = pd.DataFrame()
    df_final["Data"] = df_d["Data"].dt.strftime("%d/%m/%Y")
    for h in headers_horas:
        df_final[h] = pd.to_numeric(df_d[h], errors="coerce").fillna(0.0)

    totais = ["TOTAIS:"]
    for col in df_final.columns[1:]:
        totais.append(df_final[col].sum())

    if format_hhmm:
        for col in df_final.columns[1:]:
            df_final[col] = df_final[col].apply(_decimal_to_hhmm)
        totais = [totais[0]] + [_decimal_to_hhmm(x) for x in totais[1:]]

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), rightMargin=20, leftMargin=20, topMargin=20, bottomMargin=12)
    styles = getSampleStyleSheet()

    header_cells = ["Data"] + headers_horas
    
    pdf_rows = [header_cells]
    for _, row in df_final.iterrows():
        pdf_row = [row.iloc[0]]
        for val in row.iloc[1:]:
            if isinstance(val, (int, float)):
                pdf_row.append(f"{val:.2f}".replace(".", ","))
            else:
                pdf_row.append(str(val))
        pdf_rows.append(pdf_row)

    pdf_tot = [totais[0]]
    for val in totais[1:]:
        if isinstance(val, (int, float)):
            pdf_tot.append(f"{val:.2f}".replace(".", ","))
        else:
            pdf_tot.append(str(val))
    pdf_rows.append(pdf_tot)

    col_widths = [1.8 * cm] + [1.36 * cm] * (len(header_cells) - 1)
    tabela = Table(pdf_rows, colWidths=col_widths, repeatRows=1)

    estilo = TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#183C5F")),
        ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
    ])

    estilo.add("BACKGROUND", (1, 1), (-1, -1), colors.HexColor("#FFDAB9"))

    for r in range(1, len(pdf_rows)):
        for c in range(1, len(pdf_rows[0])):
            val_str = pdf_rows[r][c]
            if val_str.startswith("-") or (not format_hhmm and _is_negative_float(val_str)):
                estilo.add("TEXTCOLOR", (c, r), (c, r), colors.red)
                estilo.add("FONTNAME", (c, r), (c, r), "Helvetica-Bold")

    tabela.setStyle(estilo)
    story = [
        Paragraph("<b>Relatório – Diferenças (Boletim − Ponto)</b>", styles["Title"]),
        Paragraph(f"<b>Funcionário:</b> {employee}", styles["Normal"]),
        Paragraph(f"<b>Período:</b> {dt_ini:%d/%m/%Y} a {dt_fim:%d/%m/%Y}", styles["Normal"]),
        Spacer(1, 8),
        tabela,
    ]
    doc.build(story)
    return buf.getvalue()


# ============================================================
# EXPORTADORES CONSOLIDADOS POR CONTRATO
# ============================================================

def exportar_totais_consolidados_zip(
    service: Any,
    lista_contratos: List[str],
    dt_ini: datetime,
    dt_fim: datetime,
    format_hhmm: bool = False
) -> bytes:
    dt_ini_str = dt_ini.strftime("%Y-%m-%d")
    dt_fim_str = dt_fim.strftime("%Y-%m-%d")
    df_boletim_filtrado = service.get_boletins_df(lista_contratos, dt_ini_str, dt_fim_str)
    todos_funcionarios = sorted(df_boletim_filtrado["Funcionário"].dropna().unique())

    resultados = []
    for nome in todos_funcionarios:
        totais = service.calcular_totais_funcionario(nome, dt_ini, dt_fim)
        if not totais:
            continue
        contratos_do_func = df_boletim_filtrado[df_boletim_filtrado["Funcionário"] == nome]["Contrato"].unique()
        totais["Contratos"] = ", ".join(map(str, contratos_do_func))
        resultados.append(totais)

    if not resultados:
        # Retorna um zip vazio ou com uma mensagem simples
        zip_buf = io.BytesIO()
        with zipfile.ZipFile(zip_buf, "w") as zf:
            zf.writestr("sem_dados.txt", "Nenhum dado encontrado para o período/contratos selecionados.")
        return zip_buf.getvalue()

    df_all = pd.DataFrame(resultados)

    # ------------------ WORKBOOK 1: TOTAIS CONSOLIDADOS ------------------
    wb_tot = Workbook()
    wb_tot.remove(wb_tot.active) # Remove default sheet
    used_names = set()

    # Abas por contrato
    for contrato_id in sorted(set(lista_contratos)):
        sheet_name = _sheet_name_unique(contrato_id, used_names)
        df_contrato = df_all[df_all["Contratos"].apply(lambda x: str(contrato_id) in str(x))].copy()
        if df_contrato.empty:
            continue

        df_contrato.drop(columns=["Contratos"], inplace=True, errors="ignore")
        if "Funcionário" in df_contrato.columns:
            df_contrato = df_contrato.sort_values("Funcionário", kind="stable").reset_index(drop=True)

        # Calcula linha de totais
        num_cols = list(df_contrato.select_dtypes(include="number").columns)
        linha_total = {c: 0.0 for c in num_cols}
        for c in num_cols:
            linha_total[c] = df_contrato[c].sum()
        linha_total["Funcionário"] = "TOTAIS"

        # Converte para DataFrame final com formatação se HH:MM
        df_contrato_f = df_contrato.copy()
        if format_hhmm:
            for c in num_cols:
                df_contrato_f[c] = df_contrato_f[c].apply(_decimal_to_hhmm)
            linha_total_f = {c: (_decimal_to_hhmm(linha_total[c]) if c in num_cols else linha_total.get(c, "")) for c in df_contrato.columns}
        else:
            linha_total_f = {c: (round(linha_total[c], 2) if c in num_cols else linha_total.get(c, "")) for c in df_contrato.columns}

        df_contrato_f = pd.concat([df_contrato_f, pd.DataFrame([linha_total_f])], ignore_index=True)
        df_contrato_f.columns = [str(c).replace("_", " ") for c in df_contrato_f.columns]

        ws = wb_tot.create_sheet(title=sheet_name)
        
        # Escreve dados da planilha (headers em row 6)
        startrow = 6
        for col_idx, col_name in enumerate(df_contrato_f.columns, start=1):
            ws.cell(row=startrow, column=col_idx, value=col_name)
        
        for row_idx, (_, row) in enumerate(df_contrato_f.iterrows(), start=startrow + 1):
            for col_idx, col_name in enumerate(df_contrato_f.columns, start=1):
                ws.cell(row=row_idx, column=col_idx, value=row[col_name])

        # Determina boletins incluídos
        df_b = df_boletim_filtrado[df_boletim_filtrado["Contrato"] == contrato_id].copy()
        if "BOLETIM" in df_b.columns:
            nums = (
                pd.to_numeric(df_b["BOLETIM"], errors="coerce")
                .dropna().astype(int).sort_values().unique().tolist()
            )
            boletins_txt = ", ".join(map(str, nums)) if nums else "-"
        else:
            boletins_txt = "-"

        # Aplica cabeçalho e formatação profissional
        aplicar_cabecalho_profissional(
            ws,
            titulo="Relatório Consolidado de Totais",
            contrato=contrato_id,
            dt_ini=dt_ini,
            dt_fim=dt_fim,
            boletins=boletins_txt,
            df_cols=list(df_contrato_f.columns),
        )
        formatar_planilha_profissional(ws, startrow=startrow)

    # Aba "Resumo Geral"
    df_resumo = df_all.copy()
    df_resumo["Contrato"] = df_resumo.get("Contratos", "").apply(_ult5_contratos)

    boletins_por_func = {}
    df_b_all = df_boletim_filtrado.copy()
    if "BOLETIM" in df_b_all.columns:
        for nome in df_resumo["Funcionário"].dropna().unique():
            nums = (
                pd.to_numeric(df_b_all.loc[df_b_all["Funcionário"] == nome, "BOLETIM"], errors="coerce")
                .dropna().astype(int).sort_values().unique().tolist()
            )
            boletins_por_func[nome] = ", ".join(map(str, nums)) if nums else ""
    df_resumo["Boletim"] = df_resumo["Funcionário"].map(boletins_por_func).fillna("")

    cols_rest = [c for c in df_resumo.columns if c not in ("Funcionário", "Contratos", "Contrato", "Boletim")]
    df_resumo = df_resumo[["Funcionário", "Contrato", "Boletim"] + cols_rest]
    if "Funcionário" in df_resumo.columns:
        df_resumo = df_resumo.sort_values("Funcionário", kind="stable").reset_index(drop=True)

    # Linha TOTAIS GERAIS
    num_cols_res = list(df_resumo.select_dtypes(include="number").columns)
    linha_total_res = {c: 0.0 for c in num_cols_res}
    for c in num_cols_res:
        linha_total_res[c] = df_resumo[c].sum()
    linha_total_res["Funcionário"] = "TOTAIS GERAIS"
    linha_total_res["Contrato"] = ""
    linha_total_res["Boletim"] = ""

    df_resumo_f = df_resumo.copy()
    if format_hhmm:
        for c in num_cols_res:
            df_resumo_f[c] = df_resumo_f[c].apply(_decimal_to_hhmm)
        linha_total_res_f = {c: (_decimal_to_hhmm(linha_total_res[c]) if c in num_cols_res else linha_total_res.get(c, "")) for c in df_resumo.columns}
    else:
        linha_total_res_f = {c: (round(linha_total_res[c], 2) if c in num_cols_res else linha_total_res.get(c, "")) for c in df_resumo.columns}

    df_resumo_f = pd.concat([df_resumo_f, pd.DataFrame([linha_total_res_f])], ignore_index=True)
    df_resumo_f.columns = [str(c).replace("_", " ") for c in df_resumo_f.columns]

    ws_res = wb_tot.create_sheet(title="Resumo Geral", index=0)
    
    startrow = 6
    for col_idx, col_name in enumerate(df_resumo_f.columns, start=1):
        ws_res.cell(row=startrow, column=col_idx, value=col_name)
    for row_idx, (_, row) in enumerate(df_resumo_f.iterrows(), start=startrow + 1):
        for col_idx, col_name in enumerate(df_resumo_f.columns, start=1):
            ws_res.cell(row=row_idx, column=col_idx, value=row[col_name])

    if "BOLETIM" in df_b_all.columns:
        all_nums = pd.to_numeric(df_b_all["BOLETIM"], errors="coerce").dropna().astype(int).sort_values().unique().tolist()
        boletins_all_text = ", ".join(map(str, all_nums)) if all_nums else "-"
    else:
        boletins_all_text = "-"

    aplicar_cabecalho_profissional(
        ws_res,
        titulo="Resumo Geral - Totais por Funcionário",
        contrato=", ".join(map(str, lista_contratos)),
        dt_ini=dt_ini,
        dt_fim=dt_fim,
        boletins=boletins_all_text,
        df_cols=list(df_resumo_f.columns)
    )
    formatar_planilha_profissional(ws_res, startrow=startrow)


    # ------------------ WORKBOOK 2: DIFERENÇAS ------------------
    wb_dif = Workbook()
    wb_dif.remove(wb_dif.active)
    used_d = set()

    diff_cols_raw = [c for c in df_all.columns if str(c).startswith("Diferença ")]

    if diff_cols_raw:
        for contrato_id in sorted(set(lista_contratos)):
            sheet_name = _sheet_name_unique(contrato_id, used_d)
            df_contrato = df_all[df_all["Contratos"].apply(lambda x: str(contrato_id) in str(x))].copy()
            if df_contrato.empty:
                continue

            df_contrato_dif = df_contrato[["Funcionário"] + diff_cols_raw].copy()
            if "Funcionário" in df_contrato_dif.columns:
                df_contrato_dif = df_contrato_dif.sort_values("Funcionário", kind="stable").reset_index(drop=True)

            num_cols = list(df_contrato_dif.select_dtypes(include="number").columns)
            linha_total = {c: 0.0 for c in num_cols}
            for c in num_cols:
                linha_total[c] = df_contrato_dif[c].sum()
            linha_total["Funcionário"] = "TOTAIS"

            df_contrato_dif_f = df_contrato_dif.copy()
            if format_hhmm:
                for c in num_cols:
                    df_contrato_dif_f[c] = df_contrato_dif_f[c].apply(_decimal_to_hhmm)
                linha_total_f = {c: (_decimal_to_hhmm(linha_total[c]) if c in num_cols else linha_total.get(c, "")) for c in df_contrato_dif.columns}
            else:
                linha_total_f = {c: (round(linha_total[c], 2) if c in num_cols else linha_total.get(c, "")) for c in df_contrato_dif.columns}

            df_contrato_dif_f = pd.concat([df_contrato_dif_f, pd.DataFrame([linha_total_f])], ignore_index=True)
            df_contrato_dif_f.columns = [str(c).replace("_", " ") for c in df_contrato_dif_f.columns]

            ws = wb_dif.create_sheet(title=sheet_name)

            startrow = 6
            for col_idx, col_name in enumerate(df_contrato_dif_f.columns, start=1):
                ws.cell(row=startrow, column=col_idx, value=col_name)
            for row_idx, (_, row) in enumerate(df_contrato_dif_f.iterrows(), start=startrow + 1):
                for col_idx, col_name in enumerate(df_contrato_dif_f.columns, start=1):
                    ws.cell(row=row_idx, column=col_idx, value=row[col_name])

            aplicar_cabecalho_profissional(
                ws,
                titulo="Relatório – DIFERENÇAS (Boletim − Ponto)",
                contrato=contrato_id,
                dt_ini=dt_ini,
                dt_fim=dt_fim,
                boletins="-",
                df_cols=list(df_contrato_dif_f.columns),
            )
            formatar_planilha_profissional(ws, startrow=startrow)

            # Destaca negativos
            max_row, max_col = ws.max_row, ws.max_column
            for r in range(startrow + 1, max_row + 1):
                for c in range(2, max_col + 1):
                    cell = ws.cell(row=r, column=c)
                    if eh_negativo_str(cell.value):
                        cell.font = Font(name='Calibri', size=10, bold=True, color=CORES['erro'])

        # Resumo Geral Diferenças
        df_res_dif = df_all.copy()
        df_res_dif["Contrato"] = df_res_dif.get("Contratos", "").apply(_ult5_contratos)
        keep = ["Funcionário", "Contrato"] + diff_cols_raw
        df_res_dif = df_res_dif[keep]
        if "Funcionário" in df_res_dif.columns:
            df_res_dif = df_res_dif.sort_values("Funcionário", kind="stable").reset_index(drop=True)

        num_cols_res = list(df_res_dif.select_dtypes(include="number").columns)
        linha_total_res = {c: 0.0 for c in num_cols_res}
        for c in num_cols_res:
            linha_total_res[c] = df_res_dif[c].sum()
        linha_total_res["Funcionário"] = "TOTAIS GERAIS"
        linha_total_res["Contrato"] = ""

        df_res_dif_f = df_res_dif.copy()
        if format_hhmm:
            for c in num_cols_res:
                df_res_dif_f[c] = df_res_dif_f[c].apply(_decimal_to_hhmm)
            linha_total_res_f = {c: (_decimal_to_hhmm(linha_total_res[c]) if c in num_cols_res else linha_total_res.get(c, "")) for c in df_res_dif.columns}
        else:
            linha_total_res_f = {c: (round(linha_total_res[c], 2) if c in num_cols_res else linha_total_res.get(c, "")) for c in df_res_dif.columns}

        df_res_dif_f = pd.concat([df_res_dif_f, pd.DataFrame([linha_total_res_f])], ignore_index=True)
        df_res_dif_f.columns = [str(c).replace("_", " ") for c in df_res_dif_f.columns]

        ws_res_dif = wb_dif.create_sheet(title="Resumo Geral", index=0)
        
        startrow = 6
        for col_idx, col_name in enumerate(df_res_dif_f.columns, start=1):
            ws_res_dif.cell(row=startrow, column=col_idx, value=col_name)
        for row_idx, (_, row) in enumerate(df_res_dif_f.iterrows(), start=startrow + 1):
            for col_idx, col_name in enumerate(df_res_dif_f.columns, start=1):
                ws_res_dif.cell(row=row_idx, column=col_idx, value=row[col_name])

        aplicar_cabecalho_profissional(
            ws_res_dif,
            titulo="Resumo Geral – DIFERENÇAS (Boletim − Ponto)",
            contrato=", ".join(map(str, lista_contratos)),
            dt_ini=dt_ini,
            dt_fim=dt_fim,
            boletins="-",
            df_cols=list(df_res_dif_f.columns),
        )
        formatar_planilha_profissional(ws_res_dif, startrow=startrow)

        max_row, max_col = ws_res_dif.max_row, ws_res_dif.max_column
        for r in range(startrow + 1, max_row + 1):
            for c in range(3, max_col + 1):
                cell = ws_res_dif.cell(row=r, column=c)
                if eh_negativo_str(cell.value):
                    cell.font = Font(name='Calibri', size=10, bold=True, color=CORES['erro'])

    # Salva workbooks em BytesIO
    buf_tot = io.BytesIO()
    wb_tot.save(buf_tot)
    tot_bytes = buf_tot.getvalue()

    dif_bytes = b""
    if diff_cols_raw:
        buf_dif = io.BytesIO()
        wb_dif.save(buf_dif)
        dif_bytes = buf_dif.getvalue()

    # Zippa os arquivos e retorna
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("Totais_Consolidados_por_Contrato.xlsx", tot_bytes)
        if dif_bytes:
            zf.writestr("Totais_Consolidados_por_Contrato_Diferencas.xlsx", dif_bytes)

    return zip_buffer.getvalue()


def exportar_totais_separados_por_contrato_zip(
    service: Any,
    lista_contratos: List[str],
    dt_ini: datetime,
    dt_fim: datetime,
    format_hhmm: bool = False,
) -> bytes:
    """Gera um par de planilhas (totais e diferenças) para cada contrato."""
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as destination:
        for contrato in sorted({str(c).strip() for c in lista_contratos if str(c).strip()}):
            payload = exportar_totais_consolidados_zip(
                service, [contrato], dt_ini, dt_fim, format_hhmm
            )
            safe_contract = re.sub(r"[^A-Za-z0-9_-]+", "_", contrato).strip("_") or "contrato"
            with zipfile.ZipFile(io.BytesIO(payload), "r") as source:
                for member in source.infolist():
                    if member.is_dir():
                        continue
                    original = Path(member.filename)
                    if original.suffix.lower() == ".xlsx":
                        filename = f"{original.stem}_{safe_contract}{original.suffix}"
                    else:
                        filename = f"{safe_contract}_{original.name}"
                    destination.writestr(filename, source.read(member))
        if not destination.namelist():
            destination.writestr(
                "sem_dados.txt",
                "Nenhum dado encontrado para o período/contratos selecionados.",
            )
    return output.getvalue()

def eh_negativo_str(val) -> bool:
    if val is None:
        return False
    s = str(val).strip()
    if not s:
        return False
    if re.match(r"^-?\d{1,3}:\d{2}$", s):
        return s.startswith("-")
    s2 = s.replace(".", "").replace(",", ".")
    try:
        return float(s2) < 0
    except Exception:
        return s.startswith("-")
